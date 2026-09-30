"""
Reading a tenant's student list — typed into the portal form or uploaded as
a CSV / Excel sheet — into clean Student field values.

Both paths go through normalize_row(), so a typed row and an uploaded row
are validated the same way. Every problem is reported against its row
(and S.N. when there is one) so the tenant can fix the sheet and re-upload;
the caller saves nothing while any row has an error.
"""
import csv
import io
import re
from datetime import date, datetime, time

from django.db.models import Q

from .models import Student, VehicleType, normalize_plate

MAX_ROWS = 1000

# Canonical field -> accepted header spellings (compared lowercased with
# everything but letters and digits removed, so "Class Time (From)" and
# "class_time_from" are the same header).
HEADER_ALIASES = {
    'sn': ['sn', 'sno', 'snno', 'serialno', 'serialnumber', 'serial'],
    'batch_start_date': ['batchstartdate', 'batchstart', 'startdate', 'start'],
    'batch_end_date': ['batchenddate', 'batchend', 'enddate', 'end', 'expirydate', 'validuntil'],
    'class_days': ['classdays', 'days', 'classday'],
    'class_time_from': ['classtimefrom', 'timefrom', 'from', 'classfrom', 'fromtime'],
    'class_time_to': ['classtimeto', 'timeto', 'to', 'classto', 'totime'],
    'name': ['studentname', 'name', 'fullname'],
    'contact_number': ['contactnumber', 'contact', 'contactno', 'phone', 'phoneno', 'mobile', 'mobileno'],
    'license_plate': ['vehiclenumber', 'vehicleno', 'vehicle', 'licenseplate', 'plate', 'numberplate', 'bikeno'],
    'vehicle_type': ['vehicletype', 'type'],
    'status': ['status', 'active'],
}
_HEADER_LOOKUP = {alias: field for field, aliases in HEADER_ALIASES.items() for alias in aliases}

TEMPLATE_HEADERS = [
    'SN', 'Batch Start Date', 'Batch End Date', 'Class Days', 'Class Time From',
    'Class Time To', 'Student Name', 'Contact Number', 'Vehicle Number', 'Vehicle Type', 'Status',
]
TEMPLATE_EXAMPLE = [
    '1', '2026-10-01', '2026-12-31', 'Sun-Fri', '07:00', '09:00',
    'Ram Sharma', '9800000000', 'BA 2 PA 1234', 'Bike', 'Active',
]

_DAY_NAMES = {
    'mon': 'MON', 'monday': 'MON', 'tue': 'TUE', 'tues': 'TUE', 'tuesday': 'TUE',
    'wed': 'WED', 'wednesday': 'WED', 'thu': 'THU', 'thur': 'THU', 'thurs': 'THU',
    'thursday': 'THU', 'fri': 'FRI', 'friday': 'FRI', 'sat': 'SAT', 'saturday': 'SAT',
    'sun': 'SUN', 'sunday': 'SUN',
}
_DATE_FORMATS = ['%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%m/%d/%Y', '%d %b %Y', '%d %B %Y']
_TIME_FORMATS = ['%H:%M', '%H:%M:%S', '%I:%M %p', '%I:%M%p', '%I %p', '%I%p', '%H.%M']


class StudentFileError(Exception):
    """The file itself can't be read (wrong type, no header row, too big)."""


def _header_key(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())


def _blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def parse_date(value):
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'"{text}" is not a date (use YYYY-MM-DD).')


def parse_time(value):
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    text = re.sub(r'\s+', ' ', str(value).strip().upper()).replace('.', ':')
    for fmt in _TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    raise ValueError(f'"{value}" is not a time (use e.g. 07:00 or 7:00 AM).')


def parse_days(value):
    """'Sun-Fri', 'Mon, Wed, Fri', 'Daily', or a list -> ['SUN', 'MON', ...]
    in week order."""
    if _blank(value):
        return []
    if isinstance(value, (list, tuple)):
        parts = [str(p) for p in value]
    else:
        text = str(value).strip().lower()
        if text in ('daily', 'everyday', 'every day', 'all', 'all days'):
            return list(Student.DAYS)
        parts = re.split(r'[,/;&]|\band\b', text)

    days = set()
    for part in parts:
        part = part.strip().lower()
        if not part:
            continue
        if '-' in part or ' to ' in part:
            start, end = [p.strip() for p in re.split(r'-|\bto\b', part, maxsplit=1)]
            if start not in _DAY_NAMES or end not in _DAY_NAMES:
                raise ValueError(f'"{part}" is not a day range (e.g. Sun-Fri).')
            i, j = Student.DAYS.index(_DAY_NAMES[start]), Student.DAYS.index(_DAY_NAMES[end])
            span = range(i, j + 1) if i <= j else [*range(i, 7), *range(0, j + 1)]
            days.update(Student.DAYS[k] for k in span)
        elif part in _DAY_NAMES:
            days.add(_DAY_NAMES[part])
        else:
            raise ValueError(f'"{part}" is not a day (use Sun, Mon, … or Sun-Fri).')
    return [d for d in Student.DAYS if d in days]


def parse_status(value):
    if _blank(value) or value is True:
        return True
    if value is False:
        return False
    text = str(value).strip().lower()
    if text in ('active', 'yes', 'y', 'true', '1'):
        return True
    if text in ('inactive', 'in-active', 'no', 'n', 'false', '0'):
        return False
    raise ValueError(f'"{value}" is not a status (use Active or Inactive).')


def _text(value):
    if _blank(value):
        return ''
    # Excel stores phone numbers and S.N. as floats: 9800000000.0 -> "9800000000".
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _resolve_vehicle_type(value, cache):
    """Match a Vehicle Type column by name, then by category (Car/Bike)."""
    text = _text(value)
    if not text:
        return None
    key = text.lower()
    if key not in cache:
        category = {'bike': 'BIKE', 'motorcycle': 'BIKE', 'motorbike': 'BIKE', 'scooter': 'BIKE',
                    'scooty': 'BIKE', 'car': 'CAR'}.get(key)
        query = Q(name__iexact=text)
        if category:
            query |= Q(category=category)
        cache[key] = VehicleType.objects.filter(query).order_by('is_sample', 'id').first()
    if cache[key] is None:
        raise ValueError(f'"{text}" is not a vehicle type.')
    return cache[key]


def normalize_row(raw, vehicle_type_cache=None):
    """
    One row of raw values keyed by canonical field name -> (clean, errors).
    `clean` holds Student field values; `errors` is a list of
    {'field', 'message'}. A vehicle_type given as an id (portal form) or a
    name/category (sheet) are both accepted.
    """
    cache = {} if vehicle_type_cache is None else vehicle_type_cache
    clean, errors = {}, []

    def field(name, parser):
        try:
            clean[name] = parser(raw.get(name))
        except ValueError as exc:
            errors.append({'field': name, 'message': str(exc)})

    clean['sn'] = _text(raw.get('sn'))[:20]
    clean['name'] = _text(raw.get('name'))[:150]
    clean['contact_number'] = _text(raw.get('contact_number'))[:30]
    clean['license_plate'] = normalize_plate(_text(raw.get('license_plate')))[:20]
    field('batch_start_date', parse_date)
    field('batch_end_date', parse_date)
    field('class_days', parse_days)
    field('class_time_from', parse_time)
    field('class_time_to', parse_time)
    try:
        clean['is_active'] = parse_status(raw.get('is_active', raw.get('status')))
    except ValueError as exc:
        errors.append({'field': 'status', 'message': str(exc)})

    vt = raw.get('vehicle_type')
    if isinstance(vt, int) or (isinstance(vt, str) and vt.isdigit()):
        clean['vehicle_type'] = VehicleType.objects.filter(pk=int(vt)).first()
        if clean['vehicle_type'] is None:
            errors.append({'field': 'vehicle_type', 'message': 'Unknown vehicle type.'})
    else:
        field('vehicle_type', lambda v: _resolve_vehicle_type(v, cache))

    if not clean['name']:
        errors.append({'field': 'name', 'message': 'Student name is required.'})
    if not clean['license_plate']:
        errors.append({'field': 'license_plate', 'message': 'Vehicle number is required.'})
    if 'batch_end_date' in clean and clean['batch_end_date'] is None:
        errors.append({'field': 'batch_end_date', 'message': 'Batch end date is required.'})
    start, end = clean.get('batch_start_date'), clean.get('batch_end_date')
    if start and end and end < start:
        errors.append({'field': 'batch_end_date', 'message': 'Batch end date is before the start date.'})
    t_from, t_to = clean.get('class_time_from'), clean.get('class_time_to')
    if t_from and t_to and t_to <= t_from:
        errors.append({'field': 'class_time_to', 'message': 'Class end time must be after the start time.'})
    return clean, errors


def normalize_rows(raw_rows):
    """Validate every row and flag a vehicle number repeated in the same
    list. Returns (clean_rows, errors) with errors as
    {'row', 'sn', 'field', 'message'} — row is 1-based within the list."""
    cache, clean_rows, errors = {}, [], []
    seen_plates = {}
    for index, raw in enumerate(raw_rows, start=1):
        clean, row_errors = normalize_row(raw, cache)
        plate = clean.get('license_plate')
        if plate:
            if plate in seen_plates:
                row_errors.append({'field': 'license_plate',
                                   'message': f'Same vehicle number as row {seen_plates[plate]}.'})
            else:
                seen_plates[plate] = index
        clean_rows.append(clean)
        errors.extend({'row': index, 'sn': clean.get('sn', ''), **e} for e in row_errors)
    return clean_rows, errors


def _rows_from_table(table):
    """[header, *rows] (any cell types) -> list of raw dicts keyed by
    canonical field. Blank rows are skipped; unknown columns are ignored."""
    table = [row for row in table if any(not _blank(cell) for cell in row)]
    if not table:
        raise StudentFileError('The file is empty.')
    header = [_HEADER_LOOKUP.get(_header_key(cell)) for cell in table[0]]
    missing = [label for field, label in (('name', 'Student Name'), ('license_plate', 'Vehicle Number'),
                                          ('batch_end_date', 'Batch End Date')) if field not in header]
    if missing:
        raise StudentFileError(f"Missing column(s): {', '.join(missing)}. Download the template for the expected headers.")
    if len(table) - 1 > MAX_ROWS:
        raise StudentFileError(f'Too many rows ({len(table) - 1}). Upload at most {MAX_ROWS} students at a time.')

    rows = []
    for cells in table[1:]:
        raw = {}
        for field, cell in zip(header, cells):
            if field and field not in raw:
                raw[field] = cell
        rows.append(raw)
    return rows


def read_upload(upload):
    """An uploaded file -> (source, raw rows). source is 'CSV' or 'EXCEL'."""
    name = (getattr(upload, 'name', '') or '').lower()
    if name.endswith('.csv'):
        data = upload.read()
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        return 'CSV', _rows_from_table(list(csv.reader(io.StringIO(text))))
    if name.endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except Exception:
            raise StudentFileError('Could not read this Excel file. Save it as .xlsx and try again.')
        sheet = workbook.worksheets[0]
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return 'EXCEL', _rows_from_table(table)
    if name.endswith('.xls'):
        raise StudentFileError('Old .xls files are not supported. Save the sheet as .xlsx or .csv.')
    raise StudentFileError('Upload a .csv or .xlsx file.')


def template_csv():
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TEMPLATE_HEADERS)
    writer.writerow(TEMPLATE_EXAMPLE)
    return buffer.getvalue()
